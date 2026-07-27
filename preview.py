"""
Load local previews so duplicate files can be compared visually.

Tries, in order:
  images → RAW/DNG → PDF pages → video frames → text → document text
  → hex / type card fallback (every file gets a visual or readable content)
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from models import ARCHIVE_EXTS, AUDIO_EXTS, DOCUMENT_EXTS, IMAGE_EXTS, VIDEO_EXTS
from quarantine import format_bytes

TEXT_EXTS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".html", ".htm",
    ".log", ".ini", ".cfg", ".py", ".js", ".css", ".rtf", ".svg",
    ".yml", ".yaml", ".toml", ".sh", ".bat", ".c", ".cpp", ".h", ".java",
}

RAW_EXTS = {".dng", ".cr2", ".nef", ".arw", ".orf", ".rw2", ".raf", ".raw"}
OFFICE_OLD = {".doc", ".xls", ".ppt"}

PREVIEW_MAX_CHARS = 6000
THUMB_MAX_SIZE = (900, 900)
HEX_PREVIEW_BYTES = 512


@dataclass
class FilePreview:
    """What the GUI shows for one file in a duplicate group."""

    path: Path
    kind: str  # "image", "text", "composite"
    info_text: str
    image: Optional[object] = None  # PIL Image
    text_content: str = ""
    error: str = ""


def build_info_text(
    path: Path,
    size: int,
    modified: float,
    category: str,
    role: str,
    display_path: Optional[str] = None,
) -> str:
    """Human-readable file details for the compare card."""
    from datetime import datetime

    when = datetime.fromtimestamp(modified).strftime("%Y-%m-%d %H:%M:%S")
    shown = display_path or str(path)
    name = Path(shown.split(" → ")[-1]).name if display_path else path.name
    suffix = Path(name).suffix.lower() or path.suffix.lower() or "(none)"
    return (
        f"Role: {role}\n"
        f"Name: {name}\n"
        f"Path: {shown}\n"
        f"Category: {category}\n"
        f"Size: {format_bytes(size)} ({size} bytes)\n"
        f"Modified: {when}\n"
        f"Type: {suffix}"
    )


def load_preview_for_info(file_info, role: str) -> FilePreview:
    """Preview a scanned FileInfo, including items stored inside zip archives."""
    info_text = build_info_text(
        Path(file_info.name),
        file_info.size,
        file_info.modified,
        file_info.category,
        role,
        display_path=file_info.display_path,
    )
    if file_info.is_inside_archive:
        try:
            from similarity import _extract_text, _open_bytes, _pil_from_info, is_document, is_photo

            # Photos / images
            if is_photo(file_info):
                img = _pil_from_info(file_info)
                if img is not None:
                    return FilePreview(
                        path=file_info.path,
                        kind="image",
                        info_text=info_text,
                        image=_as_thumb(img),
                    )
            # PDF page from zip bytes
            if file_info.suffix == ".pdf":
                try:
                    import fitz
                    from PIL import Image

                    data = _open_bytes(file_info, max_bytes=32 * 1024 * 1024)
                    doc = fitz.open(stream=data, filetype="pdf")
                    try:
                        if doc.page_count >= 1:
                            pix = doc[0].get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False)
                            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                            return FilePreview(
                                path=file_info.path,
                                kind="image",
                                info_text=info_text,
                                image=_as_thumb(img),
                            )
                    finally:
                        doc.close()
                except Exception:
                    pass
            if is_document(file_info) or file_info.suffix in TEXT_EXTS:
                text = _extract_text(file_info)
                if text.strip():
                    return FilePreview(
                        path=file_info.path,
                        kind="text",
                        info_text=info_text,
                        text_content=text[:PREVIEW_MAX_CHARS],
                    )
            data = _open_bytes(file_info)
            return _fallback_preview(
                Path(file_info.name),
                file_info.category,
                info_text,
                data=data,
                note="Inside zip archive",
            )
        except Exception as exc:  # noqa: BLE001
            return _fallback_preview(
                Path(file_info.name),
                file_info.category,
                info_text,
                note=str(exc),
                error=str(exc),
            )
    return load_preview(
        file_info.path, file_info.size, file_info.modified, file_info.category, role
    )


def _as_thumb(img):
    """Shrink a PIL image for the compare panel."""
    from PIL import Image

    if img.mode not in ("RGB", "RGBA"):
        img = img.convert("RGB")
    elif img.mode == "RGBA":
        background = Image.new("RGB", img.size, (40, 40, 40))
        background.paste(img, mask=img.split()[3])
        img = background
    img = img.copy()
    img.thumbnail(THUMB_MAX_SIZE)
    return img


def _type_card(path: Path, category: str, note: str = "", extra_lines: Optional[list[str]] = None) -> object:
    """
    Always-available visual: a card with file type + name (+ optional hex lines).
    """
    from PIL import Image, ImageDraw, ImageFont

    height = 280 if extra_lines else 240
    img = Image.new("RGB", (360, height), (45, 48, 55))
    draw = ImageDraw.Draw(img)
    draw.rectangle((10, 10, 350, height - 10), outline=(90, 120, 160), width=2)

    ext = (path.suffix.lower().lstrip(".") or "file").upper()
    title = category or "File"
    name = path.name
    if len(name) > 40:
        name = name[:37] + "..."

    try:
        font_big = ImageFont.truetype("DejaVuSans-Bold.ttf", 32)
        font_mid = ImageFont.truetype("DejaVuSans.ttf", 16)
        font_small = ImageFont.truetype("DejaVuSans.ttf", 12)
        font_mono = ImageFont.truetype("DejaVuSansMono.ttf", 11)
    except OSError:
        font_big = ImageFont.load_default()
        font_mid = font_big
        font_small = font_big
        font_mono = font_big

    draw.text((22, 22), f".{ext}" if ext != "FILE" else "FILE", fill=(220, 230, 245), font=font_big)
    draw.text((22, 70), title, fill=(160, 190, 230), font=font_mid)
    draw.text((22, 100), name, fill=(210, 210, 210), font=font_small)
    y = 130
    if note:
        draw.text((22, y), note[:55], fill=(150, 150, 150), font=font_small)
        y += 22
    if extra_lines:
        for line in extra_lines[:6]:
            draw.text((22, y), line[:48], fill=(180, 200, 170), font=font_mono)
            y += 16
    return img


def _hex_lines(data: bytes, limit: int = HEX_PREVIEW_BYTES) -> list[str]:
    chunk = data[:limit]
    lines = []
    for i in range(0, len(chunk), 16):
        part = chunk[i : i + 16]
        hex_part = " ".join(f"{b:02x}" for b in part)
        ascii_part = "".join(chr(b) if 32 <= b < 127 else "." for b in part)
        lines.append(f"{i:04x}  {hex_part:<47}  {ascii_part}")
    return lines


def _looks_like_text(data: bytes) -> bool:
    if not data:
        return True
    sample = data[:4000]
    # Reject if too many NULs
    if sample.count(0) > max(2, len(sample) // 50):
        return False
    try:
        text = sample.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = sample.decode("latin-1")
        except UnicodeDecodeError:
            return False
    printable = sum(1 for c in text if c.isprintable() or c in "\n\r\t")
    return printable / max(len(text), 1) >= 0.85


def _fallback_preview(
    path: Path,
    category: str,
    info_text: str,
    data: Optional[bytes] = None,
    note: str = "Content preview",
    error: str = "",
) -> FilePreview:
    """
    Last-resort preview: readable text if possible, else hex dump + type card image.
    Guarantees something useful for every file type.
    """
    if data is None:
        try:
            with path.open("rb") as handle:
                data = handle.read(max(HEX_PREVIEW_BYTES, PREVIEW_MAX_CHARS))
        except OSError as exc:
            card = _type_card(path, category, note=str(exc))
            return FilePreview(
                path=path, kind="image", info_text=info_text, image=card, error=str(exc)
            )

    if _looks_like_text(data):
        text = data.decode("utf-8", errors="replace")
        if len(text) > PREVIEW_MAX_CHARS:
            text = text[:PREVIEW_MAX_CHARS] + "\n\n… (preview truncated)"
        return FilePreview(path=path, kind="text", info_text=info_text, text_content=text, error=error)

    lines = _hex_lines(data)
    hex_text = (
        f"Binary preview (first {min(len(data), HEX_PREVIEW_BYTES)} bytes):\n\n"
        + "\n".join(lines)
    )
    card = _type_card(path, category, note=note, extra_lines=lines[:4])
    return FilePreview(
        path=path,
        kind="composite",
        info_text=info_text,
        text_content=hex_text,
        image=card,
        error=error,
    )


def _preview_pillow_image(path: Path):
    from PIL import Image

    with Image.open(path) as img:
        return _as_thumb(img.copy())


def _preview_raw(path: Path):
    import rawpy

    with rawpy.imread(str(path)) as raw:
        rgb = raw.postprocess(half_size=True, use_camera_wb=True)
    from PIL import Image

    return _as_thumb(Image.fromarray(rgb))


def _preview_pdf(path: Path):
    import fitz

    doc = fitz.open(path)
    try:
        if doc.page_count < 1:
            raise ValueError("PDF has no pages")
        page = doc[0]
        pix = page.get_pixmap(matrix=fitz.Matrix(1.2, 1.2), alpha=False)
        from PIL import Image

        img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
        return _as_thumb(img)
    finally:
        doc.close()


def _preview_video(path: Path):
    import imageio.v3 as iio
    from PIL import Image

    frame = iio.imread(path, index=0)
    img = Image.fromarray(frame)
    return _as_thumb(img)


def _preview_audio_waveform(path: Path):
    """Simple waveform visual for WAV (and best-effort for other audio)."""
    from PIL import Image, ImageDraw
    import numpy as np

    samples = None
    if path.suffix.lower() == ".wav":
        import wave

        with wave.open(str(path), "rb") as wf:
            frames = wf.readframes(min(wf.getnframes(), 200_000))
            width = wf.getsampwidth()
            channels = wf.getnchannels()
            if width == 1:
                data = np.frombuffer(frames, dtype=np.uint8).astype(np.int16) - 128
            elif width == 2:
                data = np.frombuffer(frames, dtype=np.int16)
            else:
                raise ValueError("Unsupported WAV sample width")
            if channels > 1:
                data = data.reshape(-1, channels).mean(axis=1)
            samples = data
    else:
        # Best-effort: treat file bytes as signed samples for a rough shape
        raw = path.read_bytes()[:100_000]
        samples = np.frombuffer(raw, dtype=np.int8).astype(np.float32)

    if samples is None or len(samples) < 8:
        raise ValueError("No audio samples")

    # Downsample for drawing
    target = 320
    step = max(1, len(samples) // target)
    reduced = samples[::step][:target]
    peak = float(np.max(np.abs(reduced))) or 1.0
    reduced = reduced / peak

    img = Image.new("RGB", (360, 200), (40, 44, 52))
    draw = ImageDraw.Draw(img)
    mid = 100
    draw.rectangle((10, 10, 350, 190), outline=(90, 120, 160), width=2)
    for x, val in enumerate(reduced):
        y = int(mid - val * 70)
        draw.line((20 + x, mid, 20 + x, y), fill=(120, 180, 220))
    draw.text((20, 12), f"Audio · {path.suffix.lower()}", fill=(200, 210, 220))
    draw.text((20, 170), path.name[:40], fill=(160, 160, 160))
    return img


def _preview_docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    parts = []
    for para in document.paragraphs[:80]:
        text = para.text.strip()
        if text:
            parts.append(text)
    raw = "\n".join(parts) if parts else "(Word document has no readable text paragraphs)"
    if len(raw) > PREVIEW_MAX_CHARS:
        raw = raw[:PREVIEW_MAX_CHARS] + "\n\n… (preview truncated)"
    return raw


def _preview_xlsx_text(path: Path) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        lines = []
        for sheet in wb.worksheets[:3]:
            lines.append(f"[Sheet: {sheet.title}]")
            for i, row in enumerate(sheet.iter_rows(values_only=True)):
                if i >= 25:
                    lines.append("…")
                    break
                cells = ["" if c is None else str(c) for c in row[:8]]
                lines.append(" | ".join(cells))
            lines.append("")
        raw = "\n".join(lines).strip() or "(Spreadsheet is empty)"
        if len(raw) > PREVIEW_MAX_CHARS:
            raw = raw[:PREVIEW_MAX_CHARS] + "\n\n… (preview truncated)"
        return raw
    finally:
        wb.close()


def _preview_pptx_text(path: Path) -> str:
    """Best-effort text from a PowerPoint file via zip XML."""
    import zipfile
    import re

    texts: list[str] = []
    with zipfile.ZipFile(path) as zf:
        slides = sorted(n for n in zf.namelist() if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        for name in slides[:12]:
            xml = zf.read(name).decode("utf-8", errors="ignore")
            bits = re.findall(r"<a:t>([^<]*)</a:t>", xml)
            if bits:
                texts.append(f"[{Path(name).name}]")
                texts.extend(bits[:40])
                texts.append("")
    raw = "\n".join(texts).strip() or "(No readable text in presentation)"
    if len(raw) > PREVIEW_MAX_CHARS:
        raw = raw[:PREVIEW_MAX_CHARS] + "\n\n… (preview truncated)"
    return raw


def _preview_text_file(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    if len(raw) > PREVIEW_MAX_CHARS:
        raw = raw[:PREVIEW_MAX_CHARS] + "\n\n… (preview truncated)"
    return raw


def load_preview(path: Path, size: int, modified: float, category: str, role: str) -> FilePreview:
    """
    Build the best available visual/text preview for a file.
    Never uploads anything — reads only from your disk.
    Always returns a preview (image, text, or hex fallback).
    """
    info = build_info_text(path, size, modified, category, role)
    if not path.exists():
        card = _type_card(path, category, note="File missing on disk")
        return FilePreview(path=path, kind="image", info_text=info, image=card, error="File missing")

    ext = path.suffix.lower()
    errors: list[str] = []

    # --- Photos / common images ---
    if ext in IMAGE_EXTS and ext not in RAW_EXTS:
        try:
            if ext in {".heic", ".heif"}:
                try:
                    from pillow_heif import register_heif_opener

                    register_heif_opener()
                except Exception:
                    pass
            return FilePreview(
                path=path, kind="image", info_text=info, image=_preview_pillow_image(path)
            )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"image: {exc}")

    # --- RAW / DNG ---
    if ext in RAW_EXTS or ext in IMAGE_EXTS:
        try:
            return FilePreview(path=path, kind="image", info_text=info, image=_preview_raw(path))
        except Exception as exc:  # noqa: BLE001
            errors.append(f"raw: {exc}")

    # --- PDF ---
    if ext == ".pdf":
        try:
            return FilePreview(path=path, kind="image", info_text=info, image=_preview_pdf(path))
        except Exception as exc:  # noqa: BLE001
            errors.append(f"pdf: {exc}")

    # --- Video (first frame) ---
    if ext in VIDEO_EXTS:
        try:
            return FilePreview(path=path, kind="image", info_text=info, image=_preview_video(path))
        except Exception as exc:  # noqa: BLE001
            errors.append(f"video: {exc}")

    # --- Audio waveform ---
    if ext in AUDIO_EXTS:
        try:
            return FilePreview(
                path=path, kind="image", info_text=info, image=_preview_audio_waveform(path)
            )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"audio: {exc}")

    # --- Plain text ---
    if ext in TEXT_EXTS:
        try:
            return FilePreview(
                path=path, kind="text", info_text=info, text_content=_preview_text_file(path)
            )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"text: {exc}")

    # --- Word ---
    if ext == ".docx":
        try:
            return FilePreview(
                path=path, kind="text", info_text=info, text_content=_preview_docx_text(path)
            )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"docx: {exc}")

    # --- Excel ---
    if ext in {".xlsx", ".xlsm"}:
        try:
            return FilePreview(
                path=path, kind="text", info_text=info, text_content=_preview_xlsx_text(path)
            )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"xlsx: {exc}")

    # --- PowerPoint ---
    if ext == ".pptx":
        try:
            return FilePreview(
                path=path, kind="text", info_text=info, text_content=_preview_pptx_text(path)
            )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"pptx: {exc}")

    # --- ZIP-like archives: show file list as text ---
    if ext in ARCHIVE_EXTS or ext == ".zip":
        try:
            import zipfile

            if zipfile.is_zipfile(path):
                with zipfile.ZipFile(path) as zf:
                    names = zf.namelist()
                listing = [f"Archive contents ({len(names)} entries):", *names[:80]]
                if len(names) > 80:
                    listing.append("…")
                return FilePreview(
                    path=path, kind="text", info_text=info, text_content="\n".join(listing)
                )
        except Exception as exc:  # noqa: BLE001
            errors.append(f"archive: {exc}")

    # --- Old Office / unknown: hex or text sniff ---
    note = "Binary / type preview"
    if ext in AUDIO_EXTS:
        note = "Audio (waveform unavailable — showing bytes)"
    elif ext in DOCUMENT_EXTS or ext in OFFICE_OLD:
        note = "Document bytes preview"
    elif errors:
        note = "Fallback content preview"

    return _fallback_preview(
        path,
        category,
        info,
        note=note,
        error="; ".join(errors) if errors else "",
    )
